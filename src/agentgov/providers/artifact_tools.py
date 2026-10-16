from __future__ import annotations

import csv
import io
import json
import sqlite3
from copy import copy
from pathlib import Path
from statistics import mean, median
from typing import Any


def _workspace_path(repo_tools: Any, session_id: str, relative_path: str) -> Path:
    path = str(relative_path or "artifact.bin")
    if Path(path).is_absolute() or ".." in Path(path).parts:
        raise PermissionError("artifact path must stay inside the Session workspace")
    return repo_tools._workspace_path(session_id, path)


def _write_bytes(repo_tools: Any, session_id: str, relative_path: str, content: bytes, artifact_type: str, mime_type: str) -> dict[str, Any]:
    target = _workspace_path(repo_tools, session_id, relative_path)
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_bytes(content)
    return {
        "status": "completed",
        "path": str(target),
        "relative_path": str(target.relative_to(repo_tools.root)),
        "bytes": len(content),
        "artifact_type": artifact_type,
        "mime_type": mime_type,
    }


def render_mermaid(repo_tools: Any, session_id: str, args: dict[str, Any]) -> dict[str, Any]:
    source = str(args.get("source") or args.get("content") or "graph TD\n  A[Start] --> B[Done]")
    name = str(args.get("name") or "diagram").replace("/", "_")
    mmd_path = f"artifacts/{name}.mmd"
    html = """<!doctype html><meta charset='utf-8'><script type='module'>
import mermaid from 'https://cdn.jsdelivr.net/npm/mermaid@11/dist/mermaid.esm.min.mjs';
mermaid.initialize({startOnLoad:true});
</script><pre class='mermaid'>""" + source + "</pre>"
    result = _write_bytes(repo_tools, session_id, mmd_path, source.encode("utf-8"), "diagram", "text/vnd.mermaid")
    html_result = _write_bytes(repo_tools, session_id, f"artifacts/{name}.html", html.encode("utf-8"), "diagram_preview", "text/html")
    result["preview_path"] = html_result["path"]
    result["source"] = source
    return result


def render_chart(repo_tools: Any, session_id: str, args: dict[str, Any]) -> dict[str, Any]:
    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
    except ImportError:
        return {"status": "unavailable", "error": "matplotlib is not installed; install agentgov[artifacts]"}
    labels = [str(item) for item in args.get("labels", [])]
    values = [float(item) for item in args.get("values", [])]
    if not labels or len(labels) != len(values):
        return {"status": "failed", "error": "labels and values must be equally sized non-empty arrays"}
    figure, axis = plt.subplots(figsize=(9, 5))
    chart_type = str(args.get("chart_type") or "bar")
    if chart_type == "line":
        axis.plot(labels, values, marker="o", color="#2563eb")
    else:
        axis.bar(labels, values, color="#0f766e")
    axis.set_title(str(args.get("title") or "Session data chart"))
    axis.grid(axis="y", alpha=0.2)
    figure.tight_layout()
    output = io.BytesIO()
    figure.savefig(output, format="png", dpi=150)
    plt.close(figure)
    name = str(args.get("name") or "chart").replace("/", "_")
    result = _write_bytes(repo_tools, session_id, f"artifacts/{name}.png", output.getvalue(), "chart", "image/png")
    result["labels"] = labels
    result["values"] = values
    return result


def generate_image(repo_tools: Any, session_id: str, args: dict[str, Any]) -> dict[str, Any]:
    """Create a small, inspectable weather illustration without an external image API."""
    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
        from matplotlib.patches import Circle, Rectangle
    except ImportError:
        return {"status": "unavailable", "error": "matplotlib is not installed; install agentgov[artifacts]"}
    title = str(args.get("title") or "Weather story illustration")
    description = str(args.get("weather") or args.get("prompt") or "clear sky").lower()
    figure, axis = plt.subplots(figsize=(8, 4.5), dpi=150)
    axis.set_xlim(0, 10)
    axis.set_ylim(0, 6)
    axis.axis("off")
    axis.set_facecolor("#dff3ff")
    figure.patch.set_facecolor("#dff3ff")
    if any(word in description for word in ("rain", "storm", "雨", "雷")):
        accent = "#244e6b"
        for x in (2.0, 4.2, 6.4, 8.6):
            axis.plot([x, x - 0.18], [3.9, 3.1], color="#477998", linewidth=2)
            axis.plot([x + 0.35, x + 0.17], [3.9, 3.1], color="#477998", linewidth=2)
    else:
        accent = "#e59a22"
        axis.add_patch(Circle((7.7, 4.25), 0.62, color="#f6bd45"))
    axis.add_patch(Rectangle((0, 0), 10, 1.55, color="#78b86b"))
    axis.add_patch(Circle((3.0, 4.2), 0.58, color="#f7f7f2"))
    axis.add_patch(Circle((3.7, 4.35), 0.78, color="#f7f7f2"))
    axis.add_patch(Circle((4.5, 4.2), 0.58, color="#f7f7f2"))
    axis.text(0.45, 5.35, title[:48], fontsize=15, weight="bold", color="#17324d")
    axis.text(0.45, 4.85, str(args.get("caption") or description)[:80], fontsize=9, color="#35566c")
    axis.text(5.0, 0.55, "Session artifact illustration", ha="center", fontsize=9, color=accent)
    output = io.BytesIO()
    figure.savefig(output, format="png", bbox_inches="tight")
    plt.close(figure)
    name = str(args.get("name") or "story-illustration").replace("/", "_")
    return _write_bytes(repo_tools, session_id, f"artifacts/{name}.png", output.getvalue(), "image", "image/png")


def write_spreadsheet(repo_tools: Any, session_id: str, args: dict[str, Any]) -> dict[str, Any]:
    try:
        from openpyxl import Workbook
    except ImportError:
        return {"status": "unavailable", "error": "openpyxl is not installed; install agentgov[artifacts]"}
    sheets = args.get("sheets") or {"Report": {"headers": ["Metric", "Value"], "rows": []}}
    if not isinstance(sheets, dict):
        return {"status": "failed", "error": "sheets must be an object keyed by worksheet name"}
    workbook = Workbook()
    first = True
    for sheet_name, payload in sheets.items():
        sheet = workbook.active if first else workbook.create_sheet()
        first = False
        sheet.title = str(sheet_name)[:31] or "Sheet"
        if isinstance(payload, list):
            rows = payload
            headers = []
        else:
            payload = payload if isinstance(payload, dict) else {}
            headers = payload.get("headers") or []
            rows = payload.get("rows") or []
        if headers:
            sheet.append([str(value) for value in headers])
        for row in rows:
            if isinstance(row, dict):
                ordered = [row.get(header) for header in headers] if headers else list(row.values())
                sheet.append(ordered)
            elif isinstance(row, list):
                sheet.append(row)
        sheet.freeze_panes = "A2" if headers else None
        for cell in sheet[1] if headers else []:
            font = copy(cell.font)
            font.bold = True
            cell.font = font
    output = io.BytesIO()
    workbook.save(output)
    name = str(args.get("name") or "report").replace("/", "_")
    return _write_bytes(repo_tools, session_id, f"artifacts/{name}.xlsx", output.getvalue(), "spreadsheet", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


def generate_presentation(repo_tools: Any, session_id: str, args: dict[str, Any]) -> dict[str, Any]:
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
    except ImportError:
        return {"status": "unavailable", "error": "python-pptx is not installed; install agentgov[artifacts]"}
    slides = args.get("slides") or [{"title": "Session brief", "bullets": ["No slide content supplied"]}]
    if not isinstance(slides, list):
        return {"status": "failed", "error": "slides must be an array"}
    presentation = Presentation()
    for item in slides:
        item = item if isinstance(item, dict) else {}
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = str(item.get("title") or "Untitled")
        frame = slide.placeholders[1].text_frame
        frame.clear()
        bullets = item.get("bullets") or []
        for index, bullet in enumerate(bullets):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.text = str(bullet)
            paragraph.font.size = Pt(20)
    output = io.BytesIO()
    presentation.save(output)
    name = str(args.get("name") or "session-brief").replace("/", "_")
    return _write_bytes(repo_tools, session_id, f"artifacts/{name}.pptx", output.getvalue(), "presentation", "application/vnd.openxmlformats-officedocument.presentationml.presentation")


def generate_document(repo_tools: Any, session_id: str, args: dict[str, Any]) -> dict[str, Any]:
    try:
        from docx import Document
    except ImportError:
        return {"status": "unavailable", "error": "python-docx is not installed; install agentgov[artifacts]"}
    document = Document()
    document.add_heading(str(args.get("title") or "Session brief"), level=0)
    sections = args.get("sections") or []
    for item in sections:
        item = item if isinstance(item, dict) else {"heading": "Section", "body": str(item)}
        document.add_heading(str(item.get("heading") or "Section"), level=1)
        document.add_paragraph(str(item.get("body") or ""))
    output = io.BytesIO()
    document.save(output)
    name = str(args.get("name") or "session-brief").replace("/", "_")
    return _write_bytes(repo_tools, session_id, f"artifacts/{name}.docx", output.getvalue(), "document", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")


def export_pdf(repo_tools: Any, session_id: str, args: dict[str, Any]) -> dict[str, Any]:
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas
    except ImportError:
        return {"status": "unavailable", "error": "reportlab is not installed; install agentgov[artifacts]"}
    text = str(args.get("content") or "")
    output = io.BytesIO()
    pdf = canvas.Canvas(output, pagesize=A4)
    y = 800
    for line in text.splitlines() or [""]:
        pdf.drawString(48, y, line[:110])
        y -= 16
        if y < 48:
            pdf.showPage()
            y = 800
    pdf.save()
    name = str(args.get("name") or "session-report").replace("/", "_")
    return _write_bytes(repo_tools, session_id, f"artifacts/{name}.pdf", output.getvalue(), "pdf", "application/pdf")


def analyze_csv(repo_tools: Any, session_id: str, args: dict[str, Any]) -> dict[str, Any]:
    text = str(args.get("content") or "")
    if not text and args.get("path"):
        path = _workspace_path(repo_tools, session_id, str(args["path"]))
        text = path.read_text(encoding="utf-8", errors="ignore")
    rows = list(csv.DictReader(io.StringIO(text)))
    numeric: dict[str, list[float]] = {}
    for row in rows:
        for key, value in row.items():
            try:
                numeric.setdefault(key, []).append(float(value))
            except (TypeError, ValueError):
                continue
    return {
        "status": "completed",
        "rows": len(rows),
        "columns": list(rows[0].keys()) if rows else [],
        "numeric_summary": {key: {"count": len(values), "mean": mean(values), "median": median(values), "min": min(values), "max": max(values)} for key, values in numeric.items()},
    }


def query_sqlite(repo_tools: Any, session_id: str, args: dict[str, Any]) -> dict[str, Any]:
    query = str(args.get("query") or "").strip()
    if not query.lower().startswith("select"):
        return {"status": "denied", "error": "sqlite_query only allows SELECT statements"}
    path = _workspace_path(repo_tools, session_id, str(args.get("path") or "knowledge.db"))
    if not path.exists():
        return {"status": "failed", "error": f"database not found: {path.name}"}
    with sqlite3.connect(path) as connection:
        connection.row_factory = sqlite3.Row
        rows = [dict(row) for row in connection.execute(query).fetchmany(int(args.get("limit", 100)))]
    return {"status": "completed", "rows": rows, "count": len(rows)}


def validate_artifact(repo_tools: Any, session_id: str, args: dict[str, Any]) -> dict[str, Any]:
    path = _workspace_path(repo_tools, session_id, str(args.get("path") or ""))
    if not path.exists():
        return {"status": "failed", "error": f"artifact not found: {path.name}"}
    return {"status": "completed", "path": str(path), "relative_path": str(path.relative_to(repo_tools.root)), "bytes": path.stat().st_size, "valid": path.stat().st_size > 0}
